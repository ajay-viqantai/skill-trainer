"""Step 1: convert every PDF / DOCX in data/raw/ to plain text.

    python scripts/01_extract_text.py

Writes data/text/<doc_id>.txt and data/manifest.csv. Files with almost no
text are flagged as 'empty_or_scanned' (they need OCR; skipped for now).
"""
import csv
import hashlib
import re
import unicodedata

from tqdm import tqdm

import _common  # noqa: F401  (sets up import path)
import config

import pymupdf
from docx import Document


def pdf_text(path):
    with pymupdf.open(path) as doc:
        return "\n".join(page.get_text("text") for page in doc)


def docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def clean(text):
    text = unicodedata.normalize("NFKC", text)
    lines = [re.sub(r"[ \t\u00a0]+", " ", ln).strip() for ln in text.splitlines()]
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def doc_id_for(path):
    rel = str(path.relative_to(config.RAW_DIR)).encode("utf-8")
    return hashlib.sha1(rel).hexdigest()[:12]


def main():
    config.TEXT_DIR.mkdir(parents=True, exist_ok=True)
    files = [p for p in config.RAW_DIR.rglob("*")
             if p.is_file() and p.suffix.lower() in {".pdf", ".docx"}]
    skipped = [p for p in config.RAW_DIR.rglob("*.doc")]
    print(f"Found {len(files)} PDF/DOCX files in {config.RAW_DIR}")
    if skipped:
        print(f"Skipping {len(skipped)} old .doc files (convert them to .docx first)")

    rows, counts = [], {"ok": 0, "empty_or_scanned": 0, "error": 0}
    for path in tqdm(files, desc="Extracting"):
        did = doc_id_for(path)
        try:
            raw = pdf_text(path) if path.suffix.lower() == ".pdf" else docx_text(path)
            text = clean(raw)
            status = "ok" if len(text) >= config.MIN_TEXT_CHARS else "empty_or_scanned"
            if status == "ok":
                (config.TEXT_DIR / f"{did}.txt").write_text(text, encoding="utf-8")
            rows.append([did, str(path.relative_to(config.RAW_DIR)), len(text), status, ""])
        except Exception as e:  # corrupted / password-protected files
            status = "error"
            rows.append([did, str(path.relative_to(config.RAW_DIR)), 0, status, str(e)[:200]])
        counts[status] += 1

    with open(config.MANIFEST, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["doc_id", "source", "chars", "status", "error"])
        w.writerows(rows)

    print(f"Done: {counts}. Report: {config.MANIFEST}")


if __name__ == "__main__":
    main()
